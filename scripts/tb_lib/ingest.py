"""Ingest study material into text, page renders and candidate figure crops.

Per source (slug) this writes:
  _src/text/<slug>.txt        text per page/slide, with =====PAGE n===== markers
  _src/render/<slug>/NNN.png  full page/slide image (to LOOK at diagrams, code and formulas)
  _src/raw_figs/<slug>/...    candidate figures cropped from the renders (bounding box of each picture)
  _src/manifest.json          slug -> file, role, units, low-text units, engine
Engines: PDF via PyMuPDF; PPT/PPTX via PowerPoint COM (Windows), else LibreOffice -> PDF, else python-pptx (text and
embedded images only, no renders); DOC/DOCX via Word COM or LibreOffice -> PDF, else python-docx (text only).
Sources with role "reference" (a published textbook used for back-filling) are ingested as text only.
"""
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import fitz
from PIL import Image


def norm_text(s):
    return re.sub(r"[ \t]+\n", "\n", s or "").strip()


def write_text(prj, slug, texts):
    prj.text.mkdir(parents=True, exist_ok=True)
    (prj.text / f"{slug}.txt").write_text("".join(f"\n=====PAGE {i}=====\n{t}\n" for i, t in texts), encoding="utf-8")


# ----------------------------------------------------------------------------- PDF
def do_pdf(prj, slug, path, figures=True, renders=True):
    cfg = prj.cfg["ingest"]
    d = fitz.open(path)
    rdir, fdir = prj.render / slug, prj.raw_figs / slug
    if renders:
        rdir.mkdir(parents=True, exist_ok=True)
    if figures:
        fdir.mkdir(parents=True, exist_ok=True)
    texts, nfig = [], 0
    for i, p in enumerate(d, 1):
        texts.append((i, norm_text(p.get_text())))
        if renders:
            k = cfg["render_width"] / p.rect.width
            p.get_pixmap(matrix=fitz.Matrix(k, k)).save(str(rdir / f"{i:03d}.png"))
        if figures:
            page_area = p.rect.width * p.rect.height
            for k, info in enumerate(p.get_image_info(xrefs=True)):
                r = fitz.Rect(info["bbox"]) & p.rect
                if r.is_empty or r.width < 50 or r.height < 35:
                    continue
                if r.width * r.height > 0.80 * page_area:      # full-page background / image-only slide
                    continue
                p.get_pixmap(clip=r, dpi=cfg["pdf_crop_dpi"], alpha=False).save(str(fdir / f"p{i:03d}_k{k}.png"))
                nfig += 1
    return texts, nfig


# ----------------------------------------------------------------- PowerPoint (COM)
def _powerpoint():
    import win32com.client
    return win32com.client.Dispatch("PowerPoint.Application")


def do_ppt_com(prj, slug, path, pp):
    cfg = prj.cfg["ingest"]
    rdir, fdir, tmp = prj.render / slug, prj.raw_figs / slug, prj.src / "_tmp"
    for x in (rdir, fdir, tmp):
        x.mkdir(parents=True, exist_ok=True)
    pres = pp.Presentations.Open(str(path), True, False, False)
    sw, sh = pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight
    texts, nfig = [], 0

    def walk(shapes):
        for s in shapes:
            try:
                if s.Type == 6:
                    yield from walk(s.GroupItems)
                    continue
            except Exception:
                pass
            yield s

    def gettext(s):
        out = []
        try:
            if s.HasTextFrame:
                out.append(s.TextFrame.TextRange.Text)
        except Exception:
            pass
        try:
            if s.HasTable:
                for r in range(1, s.Table.Rows.Count + 1):
                    out.append(" | ".join(s.Table.Cell(r, c).Shape.TextFrame.TextRange.Text
                                          for c in range(1, s.Table.Columns.Count + 1)))
        except Exception:
            pass
        return "\n".join(out)

    try:
        for i in range(1, pres.Slides.Count + 1):
            sl = pres.Slides(i)
            tf = tmp / f"{slug}_{i:03d}.png"
            sl.Export(str(tf), "PNG", cfg["crop_width"], int(round(cfg["crop_width"] * sh / sw)))
            big = Image.open(tf).convert("RGB")
            big.resize((cfg["render_width"], int(round(cfg["render_width"] * big.height / big.width))), Image.LANCZOS).save(rdir / f"{i:03d}.png")
            parts = []
            for s in walk(sl.Shapes):
                parts.append(gettext(s))
                try:
                    pic = s.Type in (13, 11, 28) or (s.Type == 14 and s.PlaceholderFormat.ContainedType == 13)
                except Exception:
                    pic = False
                if pic:
                    sc = big.width / sw
                    L, T, W, H = s.Left * sc, s.Top * sc, s.Width * sc, s.Height * sc
                    if W < 60 or H < 40 or W * H > 0.8 * big.width * big.height:
                        continue
                    box = (max(0, int(L)), max(0, int(T)), min(big.width, int(L + W)), min(big.height, int(T + H)))
                    if box[2] - box[0] < 40 or box[3] - box[1] < 30:
                        continue
                    big.crop(box).save(fdir / f"p{i:03d}_k{nfig}.png")
                    nfig += 1
            try:
                notes = "\n".join(s.TextFrame.TextRange.Text for s in sl.NotesPage.Shapes if s.HasTextFrame)
                if notes.strip():
                    parts.append("NOTES: " + notes)
            except Exception:
                pass
            texts.append((i, norm_text("\n".join(x for x in parts if x))))
            tf.unlink(missing_ok=True)
    finally:
        pres.Close()
    return texts, nfig


# ------------------------------------------------------------ office -> PDF helpers
def _soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(p).exists():
            return p
    return None


def office_to_pdf(path):
    """Convert an office document to PDF with LibreOffice; returns the PDF path or None."""
    exe = _soffice()
    if not exe:
        return None
    out = Path(tempfile.mkdtemp(prefix="tb_office_"))
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    pdfs = list(out.glob("*.pdf"))
    return pdfs[0] if pdfs else None


def word_to_pdf_com(path):
    try:
        import win32com.client
        w = win32com.client.Dispatch("Word.Application")
        out = Path(tempfile.mkdtemp(prefix="tb_word_")) / (Path(path).stem + ".pdf")
        doc = w.Documents.Open(str(path), ReadOnly=True)
        doc.ExportAsFixedFormat(str(out), 17)
        doc.Close(False)
        w.Quit()
        return out if out.exists() else None
    except Exception:
        return None


# ------------------------------------------------------------- python-pptx fallback
def do_pptx_python(prj, slug, path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    fdir = prj.raw_figs / slug
    fdir.mkdir(parents=True, exist_ok=True)
    pr = Presentation(path)
    texts, nfig = [], 0

    def walk(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(s.shapes)
            else:
                yield s

    for i, sl in enumerate(pr.slides, 1):
        parts = []
        for s in walk(sl.shapes):
            if s.has_text_frame:
                parts.append(s.text_frame.text)
            if getattr(s, "has_table", False) and s.has_table:
                for r in s.table.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    im = Image.open(__import__("io").BytesIO(s.image.blob)).convert("RGB")
                    if im.width >= 60 and im.height >= 40:
                        im.save(fdir / f"p{i:03d}_k{nfig}.png")
                        nfig += 1
                except Exception:
                    pass
        if sl.has_notes_slide:
            parts.append("NOTES: " + sl.notes_slide.notes_text_frame.text)
        texts.append((i, norm_text("\n".join(parts))))
    return texts, nfig


def do_docx_python(prj, slug, path):
    import docx
    d = docx.Document(path)
    parts, texts = [], []
    for p in d.paragraphs:
        parts.append(p.text)
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text for c in r.cells))
    # a word file has no pages: group ~40 paragraphs per pseudo-page so the page markers stay useful
    chunk = 40
    body = [x for x in parts if x.strip()]
    for k in range(0, max(1, len(body)), chunk):
        texts.append((k // chunk + 1, "\n".join(body[k:k + chunk])))
    return texts, 0


# ------------------------------------------------------------------------- driver
def ingest_one(prj, slug, spec, engine="auto", com_app=None):
    """Return (manifest entry, powerpoint app or None)."""
    path = prj.source_path(slug)
    ext = path.suffix.lower()
    role = spec.get("role", "material")
    t0 = time.time()
    used = ""
    texts, nfig, renders = [], 0, True
    if role == "reference" and ext == ".pdf":
        texts, nfig = do_pdf(prj, slug, path, figures=False, renders=False)
        used, renders = "pymupdf-text-only", False
    elif ext == ".pdf":
        texts, nfig = do_pdf(prj, slug, path)
        used = "pymupdf"
    elif ext in (".pptx", ".ppt"):
        done = False
        if engine in ("auto", "com"):
            try:
                com_app = com_app or _powerpoint()
                texts, nfig = do_ppt_com(prj, slug, path, com_app)
                used, done = "powerpoint-com", True
            except Exception as e:
                if engine == "com":
                    raise
                print(f"  (PowerPoint COM unavailable for {path.name}: {type(e).__name__}) trying LibreOffice", file=sys.stderr)
        if not done and engine in ("auto", "soffice"):
            pdf = office_to_pdf(path)
            if pdf:
                texts, nfig = do_pdf(prj, slug, pdf)
                used, done = "libreoffice-pdf", True
        if not done:
            if ext == ".ppt":
                raise RuntimeError(f"{path.name}: .ppt needs PowerPoint or LibreOffice; save it as .pptx or .pdf and retry")
            texts, nfig = do_pptx_python(prj, slug, path)
            used, renders = "python-pptx (text + embedded images only, NO renders)", False
    elif ext in (".docx", ".doc"):
        pdf = word_to_pdf_com(path) if engine in ("auto", "com") else None
        pdf = pdf or (office_to_pdf(path) if engine in ("auto", "soffice") else None)
        if pdf:
            texts, nfig = do_pdf(prj, slug, pdf)
            used = "word/libreoffice-pdf"
        elif ext == ".docx":
            texts, nfig = do_docx_python(prj, slug, path)
            used, renders = "python-docx (text only)", False
        else:
            raise RuntimeError(f"{path.name}: .doc needs Word or LibreOffice; save it as .docx or .pdf and retry")
    elif ext in (".txt", ".md", ".rst", ".tex"):
        body = path.read_text(encoding="utf-8", errors="ignore").splitlines()
        texts = [(k // 60 + 1, "\n".join(body[k:k + 60])) for k in range(0, max(1, len(body)), 60)]
        used, renders = "plain text", False
    else:
        raise RuntimeError(f"unsupported file type {ext} ({path.name}); convert it to PDF")
    write_text(prj, slug, texts)
    low = [i for i, t in texts if len(t) < 60] if renders else []
    entry = {"file": spec["file"], "role": role, "units": len(texts), "low_text_units": low, "raw_figs": nfig,
             "engine": used, "seconds": round(time.time() - t0)}
    return entry, com_app


def run(prj, slugs=None, engine="auto", force=False):
    manifest = prj.manifest()
    com_app = None
    todo = [s for s in prj.cfg["sources"] if (not slugs or s in slugs)]
    if not todo:
        raise SystemExit("no sources in book.json; run  tb.py scan  first")
    for slug in todo:
        spec = prj.cfg["sources"][slug]
        if spec.get("role") == "syllabus" and Path(spec["file"]).suffix.lower() not in (".pdf", ".docx", ".pptx", ".txt", ".md"):
            continue
        if slug in manifest and not force and not slugs:
            print(f"skip {slug} (already ingested; --force to redo)")
            continue
        try:
            entry, com_app = ingest_one(prj, slug, spec, engine, com_app)
        except Exception as e:
            print(f"FAILED {slug}: {e}")
            continue
        manifest[slug] = entry
        prj.save_manifest(manifest)
        print(f"{slug:12s} units={entry['units']:4d} low-text={len(entry['low_text_units']):3d} figs={entry['raw_figs']:3d} "
              f"{entry['seconds']:4d}s  {entry['engine']}  <- {spec['file']}", flush=True)
    if com_app is not None:
        try:
            com_app.Quit()
        except Exception:
            pass
    return manifest


def scan(prj, add=True):
    """List the files in material_dir and propose slugs; merge them into book.json."""
    from .config import slugify
    exts = {".pdf", ".pptx", ".ppt", ".docx", ".doc", ".txt", ".md"}
    # skip this project's own folder (build/, Final/, _src/, figures/) when it sits inside the material folder, and _chk_ scratch files
    root = prj.root.resolve()
    files = sorted(p for p in prj.material_dir.rglob("*") if p.is_file() and p.suffix.lower() in exts and "_src" not in p.parts
                   and root not in p.resolve().parents and not p.name.startswith("_chk_"))
    taken = set(prj.cfg["sources"])
    known = {v["file"] for v in prj.cfg["sources"].values()}
    rows = []
    for p in files:
        rel = str(p.relative_to(prj.material_dir)).replace("\\", "/")
        if rel in known:
            continue
        slug = slugify(p.name, taken)
        taken.add(slug)
        role = "material"
        low = p.name.lower()
        if re.search(r"syllab|curricul|course[-_ ]?(plan|outline|description)", low):
            role = "syllabus"
        elif p.suffix.lower() == ".pdf":
            try:
                if len(fitz.open(p)) > 350:
                    role = "reference"          # a whole published textbook
            except Exception:
                pass
        rows.append((slug, rel, role))
        if add:
            prj.cfg["sources"][slug] = {"file": rel, "role": role}
    if add and rows:
        bak = prj.root / "book.json.bak"
        if (prj.root / "book.json").exists():
            shutil.copy(prj.root / "book.json", bak)
        prj.save()
        print(f"{len(rows)} new file(s) registered (previous book.json saved as {bak.name})")
    return rows
